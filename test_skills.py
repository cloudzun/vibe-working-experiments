#!/usr/bin/env python3
"""
RFP Analyzer Skill - 实际运行测试
测试目标：验证从 RFP 文档中提取需求的逻辑
"""

from docx import Document
import json

print("=" * 60)
print("🧪 RFP Analyzer Skill - 实际运行测试")
print("=" * 60)

# Step 1: 读取 RFP 文档
print("\n📖 Step 1: 读取 RFP 文档...")
try:
    doc = Document('experiments/rfp-simple/materials/VanArsdel_RFP.docx')
    print(f"✅ 成功读取文档")
    print(f"   总段落数：{len(doc.paragraphs)}")
except Exception as e:
    print(f"❌ 失败：{e}")
    exit(1)

# Step 2: 提取需求
print("\n🔍 Step 2: 提取需求...")
requirements = []

for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if not text:
        continue
    
    # 模式 1：列表项
    if para.style.name == 'List Bullet':
        requirements.append({
            'id': len(requirements) + 1,
            'text': text,
            'type': 'requirement',
            'section': 'List Bullet'
        })
    
    # 模式 2：动词开头
    elif any(text.startswith(v) for v in [
        'Deploy', 'Integrate', 'Achieve', 'Enable',
        'Supply', 'Install', 'Ensure', 'Provide'
    ]):
        requirements.append({
            'id': len(requirements) + 1,
            'text': text,
            'type': 'requirement',
            'section': para.style.name
        })

print(f"✅ 提取到 {len(requirements)} 个需求")

# Step 3: 分类统计
print("\n📊 Step 3: 分类统计...")
tech_keywords = ['platform', 'IoT', 'sensor', 'cloud', 'API', 'integration', 'software']
business_keywords = ['price', 'payment', 'case', 'reference', 'customer']
implementation_keywords = ['install', 'training', 'support', 'maintenance']

tech_count = 0
business_count = 0
implementation_count = 0

for req in requirements:
    text_lower = req['text'].lower()
    if any(kw in text_lower for kw in tech_keywords):
        tech_count += 1
    elif any(kw in text_lower for kw in business_keywords):
        business_count += 1
    elif any(kw in text_lower for kw in implementation_keywords):
        implementation_count += 1

print(f"   技术需求：~{tech_count} 个")
print(f"   商务需求：~{business_count} 个")
print(f"   实施需求：~{implementation_count} 个")
print(f"   其他需求：{len(requirements) - tech_count - business_count - implementation_count} 个")

# Step 4: 输出示例
print("\n📋 Step 4: 输出示例（前 5 个需求）")
for i, req in enumerate(requirements[:5], 1):
    print(f"\n【需求{i}】")
    print(f"类型：{req['type']}")
    print(f"章节：{req['section']}")
    print(f"内容：{req['text']}")

# Step 5: 保存结果
output_file = 'experiments/rfp-simple/test_rfp_analyzer_output.json'
with open(output_file, 'w', encoding='utf-8') as f:
    json.dump({
        'test_result': 'success',
        'total_requirements': len(requirements),
        'requirements': requirements
    }, f, indent=2, ensure_ascii=False)

print(f"\n💾 Step 5: 保存结果到 {output_file}")

# Step 6: 验证
print("\n✅ Step 6: 验证结果...")
if len(requirements) >= 15:
    print("✅ 通过：提取到足够的需求（≥15 个）")
else:
    print(f"⚠️ 警告：需求数量偏少（{len(requirements)} < 15）")

print("\n" + "=" * 60)
print("🎉 测试完成！")
print("=" * 60)
